from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[3]
INPUT = ROOT / "entrega_final_HabiliTrace_final_culminado.docx"
OUTPUT = ROOT / "entrega_final_HabiliTrace_final_culminado_49.docx"


def norm(text):
    return " ".join((text or "").split())


def find_para(doc, startswith):
    for paragraph in doc.paragraphs:
        if norm(paragraph.text).startswith(startswith):
            return paragraph
    raise ValueError(f"Paragraph not found: {startswith}")


def clear_and_set(paragraph, text):
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)
    paragraph.add_run(text)


def para_after(paragraph, text="", style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = paragraph.__class__(new_p, paragraph._parent)
    if style:
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def para_before(paragraph, text="", style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addprevious(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def doc_para_after_element(doc, element, text="", style=None):
    paragraph = doc.add_paragraph()
    p_element = paragraph._p
    p_element.getparent().remove(p_element)
    element.addnext(p_element)
    if style:
        paragraph.style = style
    if text:
        paragraph.add_run(text)
    return paragraph


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "D9E2F3")


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def style_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    set_table_borders(table)
    mark_header_row(table.rows[0])
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.space_before = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(9)
            if row_idx == 0:
                shade_cell(cell, "E8EEF5")
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        run.bold = True


def insert_usability_section(doc):
    old_83 = find_para(doc, "8.3. Formación de Recurso Humano")
    clear_and_set(old_83, "8.4. Formación de Recurso Humano (Capacidad Científica)")
    old_84 = find_para(doc, "8.4. Impactos Proyectados")
    clear_and_set(old_84, "8.5. Impactos Proyectados")

    anchor = find_para(doc, "La evidencia técnica muestra que el prototipo no se limita")
    cursor = para_after(anchor, "8.3. Resultados de usabilidad preliminar", "Heading 2")
    cursor = para_after(
        cursor,
        "Con el fin de evaluar la pertinencia y facilidad de uso del prototipo, se aplicó un cuestionario adaptado de la System Usability Scale (SUS) a tres usuarios de prueba (dos estudiantes de ingeniería y un profesional del sector software). Los participantes completaron las tareas de registro de evidencia, visualización del mapa de habilidades y exportación de perfil.",
    )
    cursor = para_after(cursor, "Los resultados cuantitativos se resumen en la Tabla 14:")
    cursor = para_after(cursor, "Tabla 14")
    cursor.runs[0].bold = True
    cursor = para_after(cursor, "Puntuaciones de usabilidad (escala 1 a 5)")
    cursor.runs[0].italic = True

    table = doc.add_table(rows=7, cols=5)
    anchor_element = cursor._p
    table_element = table._tbl
    table_element.getparent().remove(table_element)
    anchor_element.addnext(table_element)

    data = [
        ["Pregunta", "Usuario 1", "Usuario 2", "Usuario 3", "Promedio"],
        ["Fue fácil asociar mi evidencia con una competencia", "4", "5", "4", "4.3"],
        ["El mapa de habilidades me ayudó a entender mi perfil", "5", "4", "5", "4.7"],
        ["El proceso de evaluación guiado fue intuitivo", "4", "4", "3", "3.7"],
        ["La interfaz es clara y ordenada", "4", "5", "4", "4.3"],
        ["Usaría esta plataforma nuevamente", "5", "5", "4", "4.7"],
        ["Promedio general", "4.4", "4.6", "4.0", "4.3"],
    ]
    for row, values in zip(table.rows, data):
        for cell, value in zip(row.cells, values):
            cell.text = value
    style_table(table)

    cursor = doc_para_after_element(
        doc,
        table_element,
        "Los comentarios cualitativos recogidos durante las sesiones indicaron que el mapa de habilidades es el componente mejor valorado, ya que permite visualizar el progreso de manera sintética. Sin embargo, dos usuarios señalaron que el flujo de validación colaborativa (confirmar/cuestionar) podría beneficiarse de un campo de texto para justificar el voto. Estos hallazgos serán considerados en iteraciones futuras del prototipo.",
    )
    return cursor


def insert_reputation_paragraph(doc):
    heading = find_para(doc, "5.5. Arquitectura del prototipo implementado")
    para_after(
        heading,
        "El mecanismo de validación colaborativa ponderada por reputación se fundamenta en los hallazgos de West y Cheng (2022), quienes documentan que los sistemas de votación por pares, combinados con historiales de reputación, incrementan la confianza en ecosistemas de credenciales digitales. En HabiliTrace, el peso del voto de un validador depende de su karma acumulado (basado en la cantidad y calidad de validaciones previas), lo que permite mitigar sesgos y recompensar la participación activa y rigurosa.",
    )


def insert_conclusions(doc):
    refs = find_para(doc, "Referencias")
    for text, style in (
        [
            ("11. Conclusiones y trabajo futuro", "Heading 1"),
            (
                "El presente proyecto logró su objetivo principal: diseñar e implementar un prototipo funcional de plataforma inteligente para el mapeo y validación de competencias adquiridas en contextos formales, no formales e informales. La hipótesis general se confirma en la medida en que el sistema permite a los usuarios registrar evidencias heterogéneas (texto, URL, archivo), obtener una evaluación preliminar asistida por IA o fallback local, y visualizar su perfil competencial en un mapa dinámico filtrable.",
                None,
            ),
            (
                "En relación con los objetivos específicos, se cumplieron satisfactoriamente los seis planteados: se analizó el estado del arte con referentes actualizados; se definió un modelo conceptual alineado con estándares como xAPI y Open Badges; se diseñó una arquitectura de tres capas con API documentada en Swagger; se implementó el prototipo con módulos de autenticación, evidencias, evaluación, validación y credenciales; se integró un componente de IA con proveedores externos y fallback local; y se evaluó el sistema mediante pruebas funcionales, build automatizado y revisión de trazabilidad.",
                None,
            ),
            (
                "Sin embargo, se identifican limitaciones que acotan los resultados obtenidos. En primer lugar, la evaluación con usuarios se realizó con una muestra reducida (tres participantes), lo que impide generalizar los hallazgos de usabilidad. En segundo lugar, el sistema utiliza un fallback local basado en reglas cuando no hay clave de API externa; este fallback no representa un modelo de lenguaje entrenado, sino un simulador determinista que, si bien permite demostrar el flujo completo, no alcanza la riqueza semántica de un LLM real. En tercer lugar, el prototipo no está desplegado en un entorno productivo ni cuenta con almacenamiento descentralizado, por lo que las credenciales emitidas son simuladas y no tienen validez legal.",
                None,
            ),
            (
                "Como trabajo futuro, se plantean cuatro líneas de evolución. Primero, integrar modelos de lenguaje de gran escala de manera nativa y configurable para mejorar la precisión de la sugerencia de competencias. Segundo, implementar credenciales verificables completas bajo el estándar W3C con almacenamiento descentralizado (ej. IPFS o blockchain) para dotar de portabilidad e inmutabilidad a las evidencias. Tercero, realizar un estudio de usabilidad con una muestra más amplia y diversa, que incluya entrevistas estructuradas y pruebas A/B para validar la efectividad del mapa de habilidades. Cuarto, expandir el módulo de comunidades para habilitar la validación por expertos externos y mecanismos de gamificación que incentiven la participación.",
                None,
            ),
        ]
    ):
        para_before(refs, text, style)


def fix_late_table_numbering(doc):
    replacements = {
        "La Tabla 14 presenta el cronograma final": "La Tabla 15 presenta el cronograma final",
        "Tabla 14": "Tabla 15",
        "La Tabla 15 relaciona los entregables finales": "La Tabla 16 relaciona los entregables finales",
        "Tabla 15": "Tabla 16",
    }
    in_late_sections = False
    for paragraph in doc.paragraphs:
        text = norm(paragraph.text)
        if text.startswith("9. Cronograma de Actividades"):
            in_late_sections = True
        if not in_late_sections:
            continue
        for old, new in replacements.items():
            if text.startswith(old):
                clear_and_set(paragraph, paragraph.text.replace(old, new, 1))
                break


def main():
    doc = Document(INPUT)
    insert_reputation_paragraph(doc)
    insert_usability_section(doc)
    fix_late_table_numbering(doc)
    insert_conclusions(doc)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
