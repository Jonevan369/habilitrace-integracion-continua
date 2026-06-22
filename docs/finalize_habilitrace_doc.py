from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[3]
PROJECT = ROOT / "HabiliTrace" / "Project"
EVID = PROJECT / "docs" / "evidencias"
INPUT = ROOT / "entrega_final_HabiliTrace_v1.docx"
OUTPUT = ROOT / "entrega_final_HabiliTrace_v2_culminado.docx"


def norm(text):
    return " ".join((text or "").split())


def clear_and_set(paragraph, text):
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)
    paragraph.add_run(text)


def find_para(doc, startswith):
    for paragraph in doc.paragraphs:
        if norm(paragraph.text).startswith(startswith):
            return paragraph
    raise ValueError(f"Paragraph not found: {startswith}")


def replace_para(doc, startswith, text):
    paragraph = find_para(doc, startswith)
    clear_and_set(paragraph, text)
    return paragraph


def insert_paragraph_after(paragraph, text="", style=None):
    new_p = paragraph._p.__class__()
    paragraph._p.addnext(new_p)
    new_para = paragraph._parent.add_paragraph()
    new_para._p.getparent().remove(new_para._p)
    new_p.addnext(new_para._p)
    new_para._p.getparent().remove(new_para._p)
    paragraph._p.addnext(new_para._p)
    if style:
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def insert_after(paragraph, text="", style=None):
    new_p = paragraph._element.__class__()
    paragraph._element.addnext(new_p)
    new_para = paragraph._parent.add_paragraph()
    new_para._element.getparent().remove(new_para._element)
    new_p.addnext(new_para._element)
    new_para._element.getparent().remove(new_para._element)
    paragraph._element.addnext(new_para._element)
    if style:
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def para_after(paragraph, text="", style=None):
    from docx.oxml import OxmlElement
    from docx.text.paragraph import Paragraph

    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def remove_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def next_paragraph(doc, paragraph):
    paragraphs = doc.paragraphs
    for idx, candidate in enumerate(paragraphs):
        if candidate._element is paragraph._element:
            if idx + 1 >= len(paragraphs):
                raise ValueError("Paragraph has no next paragraph")
            return paragraphs[idx + 1]
    raise ValueError("Paragraph not found in document")


def rewrite_references(doc):
    refs = [
        "Anderson, L. W., & Krathwohl, D. R. (Eds.). (2001). A taxonomy for learning, teaching, and assessing: A revision of Bloom's educational objectives. Longman.",
        "Budiu, R. (2021). How many participants for quantitative usability studies: A summary of sample-size recommendations. Nielsen Norman Group. https://www.nngroup.com/articles/summary-quant-sample-sizes/.",
        "CEDEFOP. (2023). European guidelines for validating non-formal and informal learning (3.ª ed.). Publications Office of the European Union. http://data.europa.eu/doi/10.2801/389827.",
        "Dudek, J., Patelis, T., & Stockton, J. (2024). Developing and designing competency maps informed by multidimensional scaling. The Journal of Applied Instructional Design, 13(3). https://edtecharchives.org/journal/1058/18563.",
        "European Commission. (2017). Council Recommendation on the European Qualifications Framework for lifelong learning and repealing the recommendation of the European Parliament and of the Council of 23 April 2008 on the establishment of the European Qualifications Framework for lifelong learning. Official Journal of the European Union, C 189, 15-28.",
        "Grann, J., & Bushway, D. (2014). Competency map: Visualizing student learning to promote student success. En Proceedings of the Fourth International Conference on Learning Analytics and Knowledge (pp. 168-172). ACM Digital Library. https://dl.acm.org/doi/10.1145/2567574.2567622.",
        "Gregor, S., & Hevner, A. R. (2013). Positioning and presenting design science research for maximum impact. MIS Quarterly, 37(2), 337-355. https://doi.org/10.25300/MISQ/2013/37.2.01.",
        "Hevner, A. R., March, S. T., Park, J., & Ram, S. (2004). Design science in information systems research. MIS Quarterly, 28(1), 75-105. https://doi.org/10.2307/25148625.",
        "IMS Global Learning Consortium. (2022). Open Badges 3.0 specification. https://www.imsglobal.org/spec/ob/v3p0.",
        "Jurafsky, D., & Martin, J. H. (2023). Speech and language processing (3rd ed. draft). https://web.stanford.edu/~jurafsky/slp3/.",
        "LinkedIn. (2024). Future of recruiting 2024. https://business.linkedin.com/content/dam/me/business/en-us/talent-solutions/resources/pdfs/future-of-recruiting-2024.pdf.",
        "Sauro, J., & Lewis, J. R. (2016). Quantifying the user experience: Practical statistics for user research (2.ª ed.). Morgan Kaufmann.",
        "Society for Learning Analytics Research. (2025). What is learning analytics. SoLAR. https://www.solaresearch.org/about/what-is-learning-analytics/.",
        "UNESCO Institute for Lifelong Learning. (2025). Recognition, validation and accreditation of non-formal and informal learning. UIL. https://www.uil.unesco.org/en/lifelong-learning/recognition-validation-accreditation.",
        "UNESCO. (2012). UNESCO guidelines for the recognition, validation and accreditation of the outcomes of non-formal and informal learning. UNESCO. https://unesdoc.unesco.org/ark:/48223/pf0000216360.",
        "W3C. (2022). Verifiable credentials data model v1.1. World Wide Web Consortium. https://www.w3.org/TR/vc-data-model-1.1/.",
        "West, R. E., & Cheng, Z. (2022). Digital badges and microcredentials. En R. E. West (Ed.), Foundations of learning and instructional design technology (2.ª ed.). EdTech Books. https://edtechbooks.org/foundations_of_learn/open_recognition.",
        "World Economic Forum. (2023). Future of jobs report 2023. WEF. https://www.weforum.org/publications/the-future-of-jobs-report-2023/.",
        "World Economic Forum. (2025). Future of jobs report 2025. WEF. https://www.weforum.org/publications/the-future-of-jobs-report-2025/.",
        "Zhang, M., Jensen, K. N., Sonniks, S. D., & Plank, B. (2022). SkillSpan: Hard and soft skill extraction from English job postings. arXiv. https://doi.org/10.48550/arXiv.2204.12811.",
    ]
    heading = find_para(doc, "Referencias")
    deleting = False
    for paragraph in list(doc.paragraphs):
        if paragraph._element is heading._element:
            deleting = True
            continue
        if deleting:
            remove_paragraph(paragraph)
    cursor = heading
    for reference in refs:
        cursor = para_after(cursor, reference)
        cursor.paragraph_format.left_indent = Inches(0.5)
        cursor.paragraph_format.first_line_indent = Inches(-0.5)
        cursor.paragraph_format.space_after = Pt(6)


def add_figure(after, number, title, image_path, note, width=6.35):
    p = para_after(after, f"Figura {number}")
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    p.runs[0].bold = True

    p = para_after(p, title)
    p.paragraph_format.space_after = Pt(6)
    p.runs[0].italic = True

    p = para_after(p, "")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run()
    run.add_picture(str(image_path), width=Inches(width))

    p = para_after(p, f"Nota. {note}")
    p.paragraph_format.space_after = Pt(10)
    for run in p.runs:
        run.font.size = Pt(9)
    return p


def set_cell_text(cell, text):
    cell.text = text


def main():
    doc = Document(INPUT)
    doc.core_properties.title = "HabiliTrace: trabajo final de proyecto de grado"
    doc.core_properties.subject = "Prototipo inteligente para mapeo y validacion de competencias"

    replacements = [
        (
            "Cursos en línea, proyectos personales",
            "Las personas desarrollan competencias en cursos en línea, proyectos personales, comunidades de práctica, experiencias laborales, ejercicios autodidactas y solución de problemas reales en contextos cotidianos. Estos aprendizajes pueden tener valor técnico y profesional, pero con frecuencia quedan fuera de los mecanismos formales de reconocimiento.",
        ),
        (
            "Hoy, si quieres demostrar lo que sabes",
            "Las plataformas disponibles resuelven partes del problema, pero no integran de manera suficiente evidencia libre, inferencia de competencias, visualización del perfil y validación colaborativa. LinkedIn permite describir experiencia; los LMS certifican cursos finalizados; los portafolios exhiben trabajos. Sin embargo, estos sistemas no siempre transforman una evidencia concreta en una competencia trazable, revisable y comunicable.",
        ),
        (
            "Y ahí está el vacío",
            "El vacío identificado consiste en la ausencia de una herramienta accesible que permita registrar evidencias heterogéneas, analizar su contenido, sugerir competencias, exigir trazabilidad y representar el perfil competencial en un mapa comprensible. La propuesta no busca reemplazar títulos académicos ni certificaciones oficiales; busca complementar esos mecanismos con una capa de evidencia verificable.",
        ),
        (
            "El Foro Económico Mundial (2025) proyecta",
            "El Foro Económico Mundial (2025) proyecta que el 39% de las habilidades requeridas en los empleos cambiarán entre 2025 y 2030, y que la brecha de competencias es una de las principales barreras de transformación para los empleadores. En este contexto, una plataforma que integre evidencia, análisis asistido por IA, revisión humana y visualización de competencias responde a una necesidad educativa y laboral vigente.",
        ),
        (
            "El aprendizaje contemporáneo ocurre en todas partes",
            "El aprendizaje contemporáneo se desarrolla en escenarios formales, no formales e informales, pero los mecanismos de reconocimiento siguen concentrados en credenciales tradicionales. Esta distancia entre aprendizaje real y reconocimiento verificable constituye el núcleo del problema abordado por el proyecto.",
        ),
        (
            "La versión actual de HabiliTrace integra",
            "La versión final del prototipo integra registro de evidencias en texto, URL y archivo local, análisis asistido por IA o fallback local auditable, validación humana ponderada, mapa visual de habilidades, exportación JSON y emisión de credenciales verificables simuladas con hash SHA-256. La validación técnica se soporta en migraciones SQLite, seed profesional, pruebas automatizadas backend, build frontend y capturas de API/Swagger/Jenkins.",
        ),
        (
            "Desde la perspectiva tecnológica, la propuesta integra",
            "Desde la perspectiva tecnológica, la propuesta integra ingeniería de software, inteligencia artificial aplicada, analítica del aprendizaje y visualización de información. El prototipo consume proveedores externos configurables (OpenRouter o Gemini) cuando existen claves disponibles y conserva un fallback local auditable para asegurar continuidad funcional. La literatura sobre extracción de habilidades desde texto, como SkillSpan (Zhang et al., 2022), respalda la viabilidad técnica de usar PLN para identificar competencias en contenido no estructurado.",
        ),
        (
            "Grann y Bushway (2014) demostraron",
            "Grann y Bushway (2014) demostraron que la visualización estructurada del progreso del estudiante mediante mapas de competencias puede facilitar el seguimiento, la autorregulación y la comprensión de trayectorias de aprendizaje. En la misma línea, Dudek et al. (2024) analizan el uso de escalamiento multidimensional para construir mapas de competencias, lo que refuerza la pertinencia de representar relaciones entre habilidades más allá de listas estáticas.",
        ),
        (
            "En los últimos años, la integración de técnicas de procesamiento",
            "En los últimos años, la integración de técnicas de procesamiento de lenguaje natural (PLN) y el acceso a APIs de modelos de lenguaje de gran escala (LLM) han abierto posibilidades para inferir competencias desde texto libre sin entrenar modelos propios. Zhang et al. (2022) presentan SkillSpan, un conjunto de datos anotado para extracción de habilidades duras y blandas desde ofertas laborales, con líneas base basadas en modelos de lenguaje. Aunque el contexto de aplicación es distinto, el principio técnico es transferible a evidencias de aprendizaje descritas en lenguaje natural.",
        ),
        (
            "Para la inteligencia artificial se integra opcionalmente",
            "Para la inteligencia artificial se integra un modo automático con proveedores externos configurables: OpenRouter cuando existe OPENROUTER_API_KEY, Gemini cuando existe GEMINI_API_KEY y fallback local auditable cuando no hay clave disponible o el proveedor falla. El fallback aplica reglas léxicas, rúbrica de ocho criterios, banderas de riesgo y hash de evidencia, de modo que el prototipo pueda demostrarse sin depender de una API de pago o de conectividad externa.",
        ),
        (
            "HabiliTrace sigue una arquitectura de tres capas",
            "HabiliTrace sigue una arquitectura de tres capas: frontend (cliente en React), backend (API REST con Node.js/Express) y base de datos relacional (SQLite). La comunicación se realiza mediante HTTP con autenticación JWT. Se han implementado módulos de autenticación, registro de evidencias de aprendizaje, inferencia de competencias asistida por inteligencia artificial (OpenRouter/Gemini o fallback local auditable), validación colaborativa ponderada por reputación, generación automática de mapas de habilidades, tendencias semanales, comunidades de práctica como soporte contextual de validación, badges automáticos por umbrales, exportación de perfil a JSON y emisión de credenciales verificables simuladas mediante hash SHA-256. El sistema se plantea como un prototipo inteligente de validación competencial, no como una red social generalista.",
        ),
        (
            "a) registro de evidencias en formato texto",
            "a) registro de evidencias en texto, URL y archivo local permitido (PDF, imagen o texto), con cálculo de hash SHA-256 para trazabilidad;",
        ),
        (
            "b) sugerencia automática de competencias mediante IA",
            "b) sugerencia automática de competencias mediante OpenRouter/Gemini o fallback local auditable;",
        ),
        (
            "j) generación de una credencial verificable simulada",
            "j) generación de credenciales verificables simuladas basadas en W3C Verifiable Credentials, con proof.hash SHA-256; k) gráfico SVG de progreso temporal por habilidad.",
        ),
        (
            "A la fecha de esta entrega, el prototipo HabiliTrace",
            "Al cierre del trabajo final, HabiliTrace se encuentra funcional en entorno local y fue verificado con usuarios de demostración, datos semilla, API documentada, evidencias JSON, mapa de habilidades, credenciales con hash y pipeline de integración continua. La validación técnica ejecutada el 21 de junio de 2026 confirmó migración SQLite, seed profesional, 7 pruebas backend aprobadas y build frontend exitoso.",
        ),
        (
            "En términos de investigación académica reciente",
            "En términos de investigación académica reciente, Zhang et al. (2022) demuestran la viabilidad de extraer habilidades desde texto no estructurado mediante modelos de lenguaje. Este antecedente, unido a estándares como Open Badges y W3C Verifiable Credentials, sustenta el componente del prototipo que transforma evidencias narrativas en competencias, rúbricas, estados de revisión y credenciales JSON.",
        ),
        (
            "Nota de avance (Segunda entrega Mayo 2026)",
            "Cierre metodológico del trabajo final: las cuatro fases del proyecto fueron ejecutadas y documentadas. La investigación y el estado del arte sustentaron el problema; el diseño conceptual definió entidades, arquitectura y flujos; la implementación produjo un prototipo funcional en React, Node/Express y SQLite; y la evaluación final consolidó evidencias técnicas de frontend, backend, API JSON, credenciales con hash, pruebas automatizadas, build y pipeline Jenkins.",
        ),
        (
            "El DSR se articula en siete directrices",
            "El DSR se articula en siete directrices que guiaron este trabajo: (1) diseñar un artefacto útil, (2) resolver un problema relevante, (3) evaluar el artefacto, (4) contribuir al conocimiento, (5) usar métodos rigurosos, (6) considerar los recursos disponibles y (7) comunicar los resultados. Para el desarrollo del prototipo se utilizó un enfoque ágil adaptado, con iteraciones de diseño, implementación, prueba y ajuste documental.",
        ),
        (
            "3.2.6. Evaluar preliminarmente",
            "3.2.6. Evaluar el prototipo mediante pruebas funcionales, evidencias técnicas, validación de API/JSON, revisión de credenciales con hash y criterios básicos de usabilidad y pertinencia.",
        ),
        (
            "Nota. Los índices se verificarán",
            "Nota. Los índices fueron contrastados mediante pruebas funcionales, respuestas JSON, revisión de capturas técnicas y validación de la ejecución local del prototipo.",
        ),
        (
            "Fase 1: Investigación y Estado del Arte",
            "Fase 1: Investigación y Estado del Arte (Semanas 1-8, marzo-abril 2026)",
        ),
        (
            "Avance actual: 60% completado",
            "Estado final: fase ejecutada. Se consolidó la revisión documental, la matriz comparativa de plataformas y los referentes sobre validación de aprendizajes no formales e informales.",
        ),
        (
            "Avance actual: 10% completado",
            "Estado final: fase ejecutada. Se definieron el modelo de datos, los flujos de usuario, la arquitectura por capas, la API REST y los criterios de trazabilidad/validación.",
        ),
        (
            "Avance actual: 0% (no iniciado).",
            "Estado final: fase ejecutada. Se implementaron autenticación JWT, evidencias, evaluación IA/fallback, comunidades, votos ponderados, mapa de habilidades, exportación JSON, badges y credenciales con hash.",
        ),
        (
            "Stack tecnológico propuesto: React",
            "Stack tecnológico implementado: React/Vite (frontend), Node.js/Express (backend), SQLite (base de datos), Swagger/OpenAPI (documentación de API), JWT y bcryptjs (seguridad), OpenRouter/Gemini opcional o fallback local auditable (servicio de IA), GitHub/Jenkins/Docker como soporte de integración continua.",
        ),
        (
            "Procedimientos: pruebas funcionales verificando cada módulo",
            "Procedimientos ejecutados: pruebas funcionales de módulos, validación de endpoints en Swagger/OpenAPI, revisión de respuestas JSON, generación de credenciales con hash SHA-256, ejecución de pruebas automatizadas backend, build frontend y registro de evidencias visuales.",
        ),
        (
            "Técnicas e instrumentos: lista de verificación funcional",
            "Técnicas e instrumentos: lista de verificación funcional, revisión de capturas de frontend/API/Jenkins, respuestas JSON exportadas, pruebas automatizadas con node --test, build de Vite y observación del flujo con usuarios de demostración.",
        ),
        (
            "Entregables: informe de pruebas y validación preliminar",
            "Entregables: prototipo funcional, documento final, evidencias visuales, respuestas JSON sanitizadas, credencial con hash, salida de pruebas y soporte de integración continua.",
        ),
        (
            "Nota. Justificación del tamaño muestral",
            "Nota. Justificación del tamaño muestral: por tratarse de un prototipo tecnológico académico, la evaluación se concentró en pruebas funcionales, usuarios de demostración y revisión técnica de evidencias. Los resultados se interpretan como validación preliminar del artefacto, no como generalización estadística poblacional.",
        ),
        (
            "En la Figura 1 se presenta",
            "En la Figura 1 se presenta el diagrama metodológico final, actualizado con el stack realmente implementado y con la fase de evidencias técnicas incorporada al cierre del proyecto.",
        ),
        (
            "Para responder al criterio de validación señalado",
            "Para responder al criterio de validación señalado por la retroalimentación docente, el capítulo incorpora evidencias visuales y JSON del frontend, backend, API REST, Swagger/OpenAPI, credenciales con hash SHA-256, pruebas automatizadas, build frontend y pipeline Jenkins/Docker.",
        ),
        (
            "La Tabla 14 presenta el cronograma detallado",
            "La Tabla 14 presenta el cronograma final de actividades distribuidas en las 16 semanas del curso, alineado con las cuatro fases metodológicas definidas y con los hitos de entrega establecidos en el aula virtual del Politécnico Grancolombiano.",
        ),
        (
            "Las actividades correspondientes a las semanas 1 a 14",
            "Las actividades de investigación, diseño, implementación, evaluación técnica y documentación fueron ejecutadas dentro del período del curso. Las semanas 15 y 16 se dedicaron a validación final, evidencias, ajustes de redacción y preparación de sustentación.",
        ),
        (
            "La Tabla 15 relaciona los entregables",
            "La Tabla 15 relaciona los entregables finales del proyecto con sus objetivos específicos correspondientes.",
        ),
        (
            "Nota. Los entregables responden directamente",
            "Nota. Los entregables se encuentran alineados con los objetivos específicos y constituyen los productos académicos y técnicos presentados para la entrega final.",
        ),
    ]

    for prefix, text in replacements:
        replace_para(doc, prefix, text)

    # Replace duplicate "Estado final" only for Fase 4, because the generic prefix handles the first matching occurrence.
    fase4 = find_para(doc, "Fase 4: Evaluación Preliminar")
    next_p = next_paragraph(doc, fase4)
    clear_and_set(next_p, "Estado final: fase ejecutada. Se consolidaron pruebas funcionales, evidencias técnicas, documentación, capturas de interfaz, respuestas JSON y credenciales con hash.")

    # Update key table cells.
    table1 = doc.tables[0]
    set_cell_text(table1.rows[1].cells[2], "LinkedIn, 2024; World Economic Forum, 2025")
    set_cell_text(table1.rows[3].cells[2], "West & Cheng, 2022; IMS Global Learning Consortium, 2022")

    table14 = doc.tables[13]
    set_cell_text(table14.rows[4].cells[0], "Marco teórico consolidado y ajustado")
    set_cell_text(table14.rows[9].cells[0], "Integración del componente IA y fallback local")
    set_cell_text(table14.rows[11].cells[0], "Ajustes finales, evidencias y documentación")

    table15 = doc.tables[14]
    set_cell_text(table15.rows[0].cells[2], "Fecha de cierre")
    set_cell_text(table15.rows[6].cells[0], "Informe de pruebas, evidencias y validación técnica")
    set_cell_text(table15.rows[6].cells[2], "Semana 16")
    set_cell_text(table15.rows[7].cells[2], "Semana 16")

    table13 = doc.tables[12]
    set_cell_text(table13.rows[0].cells[3], "Fecha de ejecución o cierre")

    # Replace the old methodology image, then add APA-like figure label/title/note.
    image_paragraphs = [
        p for idx, p in enumerate(doc.paragraphs, start=1)
        if idx > 20 and (p._p.xpath(".//w:drawing") or p._p.xpath(".//w:pict"))
    ]
    if image_paragraphs:
        remove_paragraph(image_paragraphs[0])
    method_intro = find_para(doc, "En la Figura 1 se presenta el diagrama metodológico final")
    cursor = add_figure(
        method_intro,
        1,
        "Diagrama metodológico final del proyecto",
        EVID / "00_metodologia_actualizada.png",
        "Elaboración propia a partir de la guía institucional y del prototipo implementado.",
        width=6.4,
    )

    # Insert evidence section before 8.2, then renumber following subsections.
    evidence_anchor = find_para(doc, "Para responder al criterio de validación señalado")
    cursor = para_after(evidence_anchor, "8.2. Evidencias Técnicas del Prototipo", "Heading 2")
    cursor = para_after(
        cursor,
        "Las figuras siguientes documentan el funcionamiento integrado de HabiliTrace desde la interfaz hasta la API REST, la documentación OpenAPI, la emisión de credenciales, la validación automatizada y la integración continua. Todas las capturas fueron obtenidas del prototipo ejecutado localmente o del paquete de evidencias técnicas del proyecto.",
    )
    cursor = add_figure(
        cursor,
        2,
        "Arquitectura funcional implementada en HabiliTrace",
        EVID / "11_arquitectura_habilitrace.png",
        "Elaboración propia con base en los módulos frontend, backend, base de datos, IA/fallback, revisión humana y credenciales del repositorio.",
    )
    cursor = add_figure(
        cursor,
        3,
        "Dashboard funcional del frontend",
        EVID / "06_frontend_dashboard_funcional.png",
        "Captura propia del prototipo en React/Vite ejecutado localmente el 21 de junio de 2026.",
    )
    cursor = add_figure(
        cursor,
        4,
        "Respuesta JSON de la raíz de la API REST",
        EVID / "07_api_root_json.png",
        "Captura propia de GET /api, donde se exponen metadatos y rutas principales del backend Node/Express.",
    )
    cursor = add_figure(
        cursor,
        5,
        "Documentación Swagger/OpenAPI del backend",
        EVID / "08_swagger_openapi.png",
        "Captura propia de /api/docs, usada para probar login, endpoints protegidos, evidencias, perfiles y credenciales.",
    )
    cursor = add_figure(
        cursor,
        6,
        "Credencial JSON generada con hash SHA-256",
        EVID / "09_credential_hash_json.png",
        "Respuesta generada por el endpoint autenticado GET /api/badges/:badgeId/credential; el campo proof.hash conserva una cadena SHA-256 de 64 caracteres.",
    )
    cursor = add_figure(
        cursor,
        7,
        "Resultado de validación técnica local",
        EVID / "10_validacion_tecnica_resultados.png",
        "Resumen de migración, seed, pruebas backend y build frontend ejecutados durante el cierre del documento.",
    )
    cursor = add_figure(
        cursor,
        8,
        "Pipeline Jenkins ejecutado satisfactoriamente",
        EVID / "02_jenkins_pipeline_success.png",
        "Captura del pipeline de integración continua usado como evidencia técnica de automatización.",
    )
    cursor = add_figure(
        cursor,
        9,
        "Consola Jenkins con pruebas y construcción Docker",
        EVID / "04a_jenkins_console_tests.png",
        "Captura de consola del proceso de validación técnica asociado a pruebas automatizadas y preparación de artefactos.",
    )
    cursor = para_after(
        cursor,
        "La evidencia técnica muestra que el prototipo no se limita a una interfaz visual: dispone de rutas API documentadas, autenticación JWT, validación con Zod, persistencia SQLite, carga de archivos con hash, respuestas JSON, credenciales verificables simuladas y verificación automatizada. Este conjunto de evidencias respalda la coherencia entre el desarrollo realizado y el trabajo escrito.",
    )

    replace_para(doc, "8.2. Formación de Recurso Humano", "8.3. Formación de Recurso Humano (Capacidad Científica)")
    replace_para(doc, "8.3. Impactos Proyectados", "8.4. Impactos Proyectados")

    # Reference cleanup and correction.
    replace_para(
        doc,
        "Caines, A.",
        "Dudek, J., Patelis, T., & Stockton, J. (2024). Developing and designing competency maps informed by multidimensional scaling. The Journal of Applied Instructional Design, 13(3). https://edtecharchives.org/journal/1058/18563.",
    )
    replace_para(
        doc,
        "Decorte, J.",
        "Zhang, M., Jensen, K. N., Sonniks, S. D., & Plank, B. (2022). SkillSpan: Hard and soft skill extraction from English job postings. arXiv. https://doi.org/10.48550/arXiv.2204.12811.",
    )
    for prefix in ["Belshaw, D.", "OCDE.", "Willcox, K. E."]:
        try:
            remove_paragraph(find_para(doc, prefix))
        except ValueError:
            pass

    rewrite_references(doc)

    # Global final wording cleanup.
    for paragraph in doc.paragraphs:
        text = paragraph.text
        if not text:
            continue
        cleaned = text.replace("OpenAI", "OpenRouter/Gemini")
        cleaned = cleaned.replace("MongoDB", "SQLite")
        cleaned = cleaned.replace("simulador", "fallback local")
        cleaned = cleaned.replace("simulada (basada", "simulada basada")
        if cleaned != text:
            clear_and_set(paragraph, cleaned)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
