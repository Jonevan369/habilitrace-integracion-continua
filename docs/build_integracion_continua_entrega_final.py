from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
EVIDENCIAS = DOCS / "evidencias"
WORKFLOW = ROOT / ".github" / "workflows" / "ci.yml"
OUT = DOCS / "Integracion continua entrega final.docx"


INK = RGBColor(20, 30, 45)
MUTED = RGBColor(88, 96, 105)
BLUE = RGBColor(21, 88, 160)
GREEN = RGBColor(28, 120, 74)
RED = RGBColor(166, 48, 48)
LIGHT_BLUE = "EAF2FB"
LIGHT_GRAY = "F6F8FA"


def set_run(run, size=10.5, bold=False, italic=False, color=INK, font="Times New Roman"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:ascii"), font)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def spacing(paragraph, before=0, after=4, line=1.05):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    node = tc_pr.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        tc_pr.append(node)
    node.set(qn("w:fill"), fill)


def cell_margins(cell, top=70, start=90, bottom=70, end=90):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, inches):
    cell.width = Inches(inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def para(doc, text="", size=10.0, bold=False, italic=False, color=INK, align=None, after=3, line=1.0):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    spacing(p, after=after, line=line)
    if text:
        r = p.add_run(text)
        set_run(r, size=size, bold=bold, italic=italic, color=color)
    return p


def heading(doc, text):
    p = para(doc, text, size=11.2, bold=True, color=BLUE, after=2)
    p.paragraph_format.keep_with_next = True
    return p


def add_metadata_table(doc):
    rows = [
        ("Proyecto", "HabiliTrace"),
        ("Asignatura", "Énfasis Profesional II - Integración Continua (código 2711)"),
        ("Entrega", "Entrega 3 - Pipeline de Integración Continua con GitHub Actions"),
        ("Tutor", "Samir Yardany Peña Rocha"),
        ("Estudiante", "John Edisson Vanegas Cuervo"),
        ("Fecha", "22 de junio de 2026"),
        ("Repositorio", "https://github.com/Jonevan369/habilitrace-integracion-continua"),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.autofit = False
    for label, value in rows:
        cells = table.add_row().cells
        set_cell_width(cells[0], 1.25)
        set_cell_width(cells[1], 5.6)
        for cell in cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_margins(cell)
        shade(cells[0], LIGHT_BLUE)
        p = cells[0].paragraphs[0]
        spacing(p, after=0, line=0.9)
        set_run(p.add_run(label), size=8.2, bold=True, color=BLUE)
        p = cells[1].paragraphs[0]
        spacing(p, after=0, line=0.9)
        set_run(p.add_run(value), size=8.2)
    para(doc, after=0)


def add_workflow_table(doc):
    workflow_text = WORKFLOW.read_text(encoding="utf-8").strip()
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    headers = table.rows[0].cells
    for idx, title in enumerate(("Parte del workflow", "Función dentro de la integración continua")):
        set_cell_width(headers[idx], 2.65 if idx == 0 else 4.2)
        shade(headers[idx], LIGHT_BLUE)
        cell_margins(headers[idx])
        p = headers[idx].paragraphs[0]
        spacing(p, after=0, line=1.0)
        set_run(p.add_run(title), size=7.7, bold=True, color=BLUE)
    rows = [
        ("name: HabiliTrace CI", "Identifica el flujo en la pestaña Actions para que el evaluador ubique la ejecución."),
        ("on: push / pull_request", "Ejecuta el pipeline automáticamente cada vez que se sube un cambio o se abre un pull request hacia main."),
        ("runs-on: ubuntu-latest", "Usa un runner Linux limpio de GitHub para validar el proyecto en un entorno reproducible."),
        ("actions/checkout@v4", "Descarga el código del repositorio dentro del runner."),
        ("actions/setup-node@v4", "Instala Node.js 20 y habilita caché de npm para dependencias."),
        ("npm ci", "Instala dependencias exactamente según package-lock.json."),
        ("npm test -w server", "Ejecuta las pruebas automatizadas del backend; si una prueba falla, el build queda rojo."),
        ("npm run build", "Compila el frontend React/Vite; si hay errores de build, el flujo falla."),
    ]
    for left, right in rows:
        cells = table.add_row().cells
        for idx, value in enumerate((left, right)):
            set_cell_width(cells[idx], 2.65 if idx == 0 else 4.2)
            cell_margins(cells[idx])
            p = cells[idx].paragraphs[0]
            spacing(p, after=0, line=1.0)
            set_run(p.add_run(value), size=7.15)
    para(doc, "Archivo .github/workflows/ci.yml:", size=8.0, bold=True, after=0, line=0.85)
    p = para(doc, workflow_text, size=5.5, color=RGBColor(36, 41, 47), after=0, line=0.74)
    for run in p.runs:
        set_run(run, size=5.5, color=RGBColor(36, 41, 47), font="Courier New")


def add_evidence(doc, image_name, title, color, note):
    image_path = EVIDENCIAS / image_name
    p = para(doc, title, size=9.5, bold=True, color=color, after=2)
    p.paragraph_format.keep_with_next = True
    if image_path.exists():
        pic = doc.add_picture(str(image_path), width=Inches(6.75))
        pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        p = para(doc, f"Pendiente insertar captura: {image_name}", size=9.5, italic=True, color=RED, after=2)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para(doc, f"Nota. {note}", size=8.2, italic=True, color=MUTED, after=3)


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)

    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")

    title = para(
        doc,
        "Integración continua entrega final",
        size=14,
        bold=True,
        color=BLUE,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=2,
    )
    title.paragraph_format.keep_with_next = True
    para(
        doc,
        "Entrega 3 - Pipeline de Integración Continua con GitHub Actions",
        size=9.5,
        italic=True,
        color=MUTED,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=5,
    )
    add_metadata_table(doc)

    heading(doc, "Concepto")
    para(
        doc,
        "La integración continua es una práctica de desarrollo en la que cada cambio subido al repositorio activa validaciones automáticas como instalación de dependencias, pruebas y compilación. Sirve para detectar errores temprano, evitar que código defectuoso llegue a la rama principal y demostrar que el proyecto sigue funcionando después de cada actualización.",
        size=9.4,
        after=3,
    )

    heading(doc, "Workflow configurado")
    add_workflow_table(doc)

    doc.add_page_break()
    heading(doc, "Evidencia de estados del build")
    add_evidence(
        doc,
        "github_actions_success.png",
        "Figura 1. Ejecución exitosa del flujo (build verde)",
        GREEN,
        "Las pruebas automatizadas pasaron y el workflow finalizó correctamente.",
    )
    doc.add_page_break()
    add_evidence(
        doc,
        "github_actions_failure.png",
        "Figura 2. Ejecución fallida del flujo (build rojo)",
        RED,
        "Se rompió una prueba de forma intencional para demostrar que GitHub Actions detecta el fallo.",
    )

    heading(doc, "Cierre")
    para(
        doc,
        "El repositorio final queda público, con el workflow de GitHub Actions en .github/workflows/ci.yml y ejecución automática ante cada push o pull request. El archivo server/.env no se publica; solo se incluye server/.env.example para documentar las variables necesarias sin exponer la API privada.",
        size=10.1,
        after=0,
    )

    doc.core_properties.title = "Integración continua entrega final"
    doc.core_properties.author = "John Edisson Vanegas Cuervo"
    doc.core_properties.subject = "Entrega 3 - GitHub Actions"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
